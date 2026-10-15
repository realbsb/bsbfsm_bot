"""Экспорт финансов перед удалением из БД."""

from __future__ import annotations

import hashlib
from collections.abc import Iterable
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from ..config import Settings
from ..logging_setup import get_logger
from ..models import (
    ExportKind,
    FinanceExport,
    Project,
    Purchase,
    PurchaseItem,
)

log = get_logger(__name__)
BOLD = Font(bold=True)


def _hash(payload: str) -> str:
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def export_accounting(project: Project, out_dir: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Смета"
    ws.append(["Проект", f"#{project.id}"])
    ws.append(["Адрес", project.address_text or ""])
    ws.append(["Заказчик", project.client_name or ""])
    ws.append(
        [
            "Дата сдачи",
            project.closed_at.strftime("%Y-%m-%d") if project.closed_at else "",
        ]
    )
    ws.append([])
    ws.append(["Показатель", "Сумма"])
    for c in ws[6]:
        c.font = BOLD
    ws.append(["Внутренняя", float(project.price_internal or 0)])
    ws.append(["Согласованная", float(project.price_agreed or 0)])
    ws.append(["Предоплата", float(project.prepay_paid or 0)])
    ws.append(["Финал", float(project.final_paid or 0)])
    path = out_dir / f"accounting_{project.id}.xlsx"
    wb.save(path)
    return path


def export_supplier(project: Project, purchases: Iterable[Purchase], out_dir: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Закуп"
    ws.append(
        [
            "Магазин",
            "Дата",
            "Наименование",
            "Характеристики",
            "Кол-во",
            "Ед.",
            "Цена",
            "Сумма",
            "Гарантия до",
        ]
    )
    for c in ws[1]:
        c.font = BOLD

    for p in purchases:
        for it in p.items:
            ws.append(
                [
                    p.shop.name if p.shop else "",
                    p.purchased_at.strftime("%Y-%m-%d") if p.purchased_at else "",
                    it.common_name,
                    it.characteristics or "",
                    float(it.qty),
                    it.unit,
                    float(it.price or 0),
                    float((it.qty or 0) * (it.price or 0)),
                    it.warranty_until.strftime("%Y-%m-%d") if it.warranty_until else "",
                ]
            )
    path = out_dir / f"supplier_{project.id}.xlsx"
    wb.save(path)
    return path


def export_1c(project: Project, out_dir: Path) -> Path:
    xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<Документ>
  <Номер>{project.id}</Номер>
  <Дата>{project.closed_at.strftime("%Y-%m-%d") if project.closed_at else ""}</Дата>
  <Контрагент>{project.client_name or ""}</Контрагент>
  <Сумма>{project.price_agreed or 0}</Сумма>
  <Предоплата>{project.prepay_paid or 0}</Предоплата>
  <ОкончательныйРасчет>{project.final_paid or 0}</ОкончательныйРасчет>
  <Адрес>{project.address_text or ""}</Адрес>
</Документ>
"""
    path = out_dir / f"1c_{project.id}.xml"
    path.write_text(xml, encoding="utf-8")
    return path


def export_book(projects: Iterable[Project], year: int, out_dir: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = f"Книга {year}"
    ws.append(
        [
            "ID",
            "Дата сдачи",
            "Заказчик",
            "Адрес",
            "Согласовано",
            "Внутренняя",
            "Предоплата",
            "Финал",
            "Маржа",
        ]
    )
    for c in ws[1]:
        c.font = BOLD
    for p in projects:
        margin = float((p.price_agreed or 0) - (p.price_internal or 0))
        ws.append(
            [
                p.id,
                p.closed_at.strftime("%Y-%m-%d") if p.closed_at else "",
                p.client_name or "",
                p.address_text or "",
                float(p.price_agreed or 0),
                float(p.price_internal or 0),
                float(p.prepay_paid or 0),
                float(p.final_paid or 0),
                margin,
            ]
        )
    path = out_dir / f"book_{year}.xlsx"
    wb.save(path)
    return path


async def clear_finance_fields(session: AsyncSession, project_id: int) -> None:
    project = await session.get(Project, project_id)
    assert project is not None
    project.price_internal = None
    project.price_agreed = None
    project.prepay_paid = None
    project.final_paid = None

    purchases = (
        (await session.execute(select(Purchase).where(Purchase.project_id == project_id)))
        .scalars()
        .all()
    )
    for p in purchases:
        await session.execute(delete(PurchaseItem).where(PurchaseItem.purchase_id == p.id))
        await session.delete(p)


async def close_project_and_export(
    session: AsyncSession,
    settings: Settings,
    project: Project,
) -> dict[str, Path]:
    assert project.closed_at is not None
    out_dir = settings.export_dir / str(project.closed_at.year)
    out_dir.mkdir(parents=True, exist_ok=True)

    purchases = (
        (await session.execute(select(Purchase).where(Purchase.project_id == project.id)))
        .scalars()
        .all()
    )
    for p in purchases:
        await session.refresh(p, attribute_names=["items", "shop"])

    acc = export_accounting(project, out_dir)
    sup = export_supplier(project, purchases, out_dir)
    one = export_1c(project, out_dir)

    year = project.closed_at.year
    year_projects = (
        (
            await session.execute(
                select(Project).where(
                    Project.closed_at.between(
                        datetime(year, 1, 1),
                        datetime(year, 12, 31),
                    )
                )
            )
        )
        .scalars()
        .all()
    )
    book = export_book(year_projects, year, out_dir)

    for kind, path in (
        (ExportKind.ACCOUNTING, acc),
        (ExportKind.SUPPLIER, sup),
        (ExportKind.ONE_C, one),
        (ExportKind.BOOK, book),
    ):
        session.add(
            FinanceExport(
                project_id=project.id,
                kind=kind,
                file_path=str(path),
                payload_hash=_hash(str(path)),
            )
        )

    await clear_finance_fields(session, project.id)
    project.archived_at = datetime.utcnow()
    project.archive_path = str(out_dir)

    log.info("finance.exported", project=project.id, dir=str(out_dir))
    return {"accounting": acc, "supplier": sup, "1c": one, "book": book}
