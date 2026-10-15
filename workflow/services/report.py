"""xlsx-отчёт по проекту."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from ..logging_setup import get_logger
from ..models import (
    Document,
    DocumentItem,
    MediaAsset,
    Project,
    ProjectStageLog,
    Stage,
    User,
)

log = get_logger(__name__)
BOLD = Font(bold=True)


async def build_project_xlsx(session: AsyncSession, project: Project) -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "Проект"
    ws.append(["Проект", f"#{project.id}"])
    ws.append(
        ["Создан", project.created_at.strftime("%Y-%m-%d %H:%M") if project.created_at else ""]
    )
    ws.append(["Закрыт", project.closed_at.strftime("%Y-%m-%d %H:%M") if project.closed_at else ""])
    ws.append(["Заказчик", project.client_name or ""])
    ws.append(["Телефон", project.client_phone or ""])
    ws.append(["Адрес", project.address_text or ""])
    ws.append(["Описание", project.description or ""])
    for row in ws["A1:B7"]:
        row[0].font = BOLD

    ws2 = wb.create_sheet("Стадии")
    ws2.append(["Этап", "Когда", "Кто", "Действие", "Причина"])
    for c in ws2[1]:
        c.font = BOLD
    logs = (
        (
            await session.execute(
                select(ProjectStageLog)
                .where(ProjectStageLog.project_id == project.id)
                .order_by(ProjectStageLog.created_at)
            )
        )
        .scalars()
        .all()
    )
    for l in logs:
        stage = await session.get(Stage, l.stage_id)
        actor = await session.get(User, l.actor_id) if l.actor_id else None
        ws2.append(
            [
                stage.code if stage else "",
                l.created_at.strftime("%Y-%m-%d %H:%M"),
                actor.first_name if actor else "",
                l.action,
                l.reason or "",
            ]
        )

    ws3 = wb.create_sheet("Документы")
    ws3.append(["Тип", "Файл", "Загружен"])
    for c in ws3[1]:
        c.font = BOLD
    docs = (
        (await session.execute(select(Document).where(Document.project_id == project.id)))
        .scalars()
        .all()
    )
    for d in docs:
        ws3.append(
            [
                d.kind.value,
                d.filename or d.tg_file_id[:32],
                d.created_at.strftime("%Y-%m-%d %H:%M"),
            ]
        )

    ws4 = wb.create_sheet("Материалы")
    ws4.append(["Наименование", "Характеристики", "Кол-во", "Ед.", "Цена"])
    for c in ws4[1]:
        c.font = BOLD
    items = (
        (
            await session.execute(
                select(DocumentItem)
                .join(Document, Document.id == DocumentItem.document_id)
                .where(Document.project_id == project.id)
            )
        )
        .scalars()
        .all()
    )
    for it in items:
        ws4.append(
            [
                it.common_name,
                it.characteristics or "",
                float(it.qty),
                it.unit,
                float(it.price) if it.price else 0,
            ]
        )

    ws5 = wb.create_sheet("Медиа")
    ws5.append(["Тип", "Финальное", "Загружено"])
    for c in ws5[1]:
        c.font = BOLD
    media = (
        (await session.execute(select(MediaAsset).where(MediaAsset.project_id == project.id)))
        .scalars()
        .all()
    )
    for m in media:
        ws5.append(
            [
                m.type.value,
                "да" if m.is_final else "нет",
                m.created_at.strftime("%Y-%m-%d %H:%M"),
            ]
        )

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    log.info("report.built", project=project.id)
    return buf.read()
